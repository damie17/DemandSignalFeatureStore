"""Fill the official Capstone One-Page Summary template with project content."""

import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.oxml import parse_xml

TEMPLATE = r"C:\Users\LKDJ\Downloads\Capstone_OnePage_Summary_Template.pptx"
prs = Presentation(TEMPLATE)

# ── Colors ──
EXL_ORANGE = RGBColor(0xFB, 0x4E, 0x0B)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
CREAM      = RGBColor(0xF3, 0xE7, 0xDF)
DARK_BG    = RGBColor(0x1C, 0x1C, 0x1C)
MUTED      = RGBColor(0x6E, 0x6A, 0x66)
HINT       = RGBColor(0x9B, 0x95, 0x8F)

def find_shape(slide, name):
    for s in slide.shapes:
        if s.name == name:
            return s
    return None

def set_text(shape, text, font_size=None, bold=None, color=None, font_name=None):
    if shape is None:
        return
    tf = shape.text_frame
    while len(tf.paragraphs) > 1:
        p = tf.paragraphs[-1]._element
        p.getparent().remove(p)
    tf.paragraphs[0].clear()
    run = tf.paragraphs[0].add_run()
    run.text = text
    if font_size:
        run.font.size = Pt(font_size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = color
    if font_name:
        run.font.name = font_name

def set_multiline(shape, lines, default_size=10.5, default_color=DARK_BG, default_name="Calibri"):
    if shape is None:
        return
    tf = shape.text_frame
    for p in list(tf.paragraphs):
        p._element.getparent().remove(p._element)
    txBody = tf._txBody
    for i, (text, size, bold, color) in enumerate(lines):
        p_xml = parse_xml(
            '<a:p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:r><a:t></a:t></a:r></a:p>'
        )
        txBody.append(p_xml)
        p = tf.paragraphs[i]
        p.runs[0].text = text
        p.runs[0].font.size = Pt(size or default_size)
        p.runs[0].font.bold = bold if bold is not None else False
        p.runs[0].font.color.rgb = color or default_color
        p.runs[0].font.name = default_name

sl = prs.slides[0]

# ── Header bar ──
# Text 2: Title
set_text(find_shape(sl, "Text 2"),
         "Demand Signal Feature Store \u2014 Architecture, Trade-offs & Limitations",
         font_size=24, bold=True, color=WHITE, font_name="Cambria")

# Text 3: Track / Domain / Team / Date
set_text(find_shape(sl, "Text 3"),
         "Track: 1   \u00b7   Domain: Supply Chain / Retail   \u00b7   Team: Damini Thandele   \u00b7   Date: 10 Oct 2026",
         font_size=9.5, color=CREAM)

# ── ARCHITECTURE section ──
# Pipeline boxes
set_text(find_shape(sl, "Text 6"),
         "Supplier\nNotes (58)", font_size=9.5, bold=True, color=WHITE)
set_text(find_shape(sl, "Text 9"),
         "GPT-4o\nExtraction", font_size=9.5, bold=True, color=WHITE)
set_text(find_shape(sl, "Text 12"),
         "Feature Store", font_size=9.5, bold=True, color=WHITE)
set_text(find_shape(sl, "Text 15"),
         "Freshness SLA\n+ PSI Drift", font_size=9.5, bold=True, color=WHITE)
set_text(find_shape(sl, "Text 18"),
         "XGBoost\nForecast", font_size=9.5, bold=True, color=WHITE)

# Key components (Text 21)
set_multiline(find_shape(sl, "Text 21"), [
    ("Model/LLM: GPT-4o via OpenRouter (extraction), XGBoost (forecasting)", 10.5, False, DARK_BG),
    ("Framework: pandas, numpy, scipy, python-dotenv", 10.5, False, DARK_BG),
    ("Data store: Feature store with lineage metadata", 10.5, False, DARK_BG),
    ("Eval/monitoring: MAPE (25.01% → 23.36%), PSI/CSI drift detection, 24h freshness SLA", 10.5, False, DARK_BG),
    ("Deployment: Jupyter notebook (single Run All), Git/GitHub", 10.5, False, DARK_BG),
])

# How it works (Text 24)
set_text(find_shape(sl, "Text 24"),
         "Synthetic ERP data (2,700 daily records, 92 POs, 5 SKUs, 4 suppliers) flows into a feature engineering pipeline that computes 15 structured features. "
         "In parallel, 58 supplier notes are sent to GPT-4o, which extracts early_delay_flag, expected_delay_days, and delay_reason into structured JSON. "
         "Both feature sets merge in a feature store with lineage tracking and a 24h freshness SLA. "
         "Two identical XGBoost models are trained — Baseline (ERP only) vs Enhanced (ERP + notes) — on a chronological 80/20 split. "
         "The enhanced model reduces MAPE from 25.01% to 23.36% (-6.6%), with early_delay_flag ranking as the #1 most important feature. "
         "The key design decision: excluding delay_reason from the model despite extracting it — PSI of 1.06 confirmed it drifts too heavily for stable prediction.",
         font_size=10.5, color=MUTED, font_name="Calibri")

# ── TRADE-OFFS table ──
for shape in sl.shapes:
    if shape.has_table:
        tbl = shape.table
        trade_offs = [
            ["LLM feature selection", "2 of 3 extracted fields",
             "Use all 3 (incl. delay_reason)", "delay_reason PSI = 1.06; too variable across LLM runs"],
            ["Model algorithm", "XGBoost (tree-based)",
             "Linear regression", "Non-linear feature interactions in supply chain data"],
            ["Hyperparameters", "Fixed (n=200, depth=5)",
             "Tuned per model", "Fair comparison; same config isolates feature value"],
        ]
        for r, row_data in enumerate(trade_offs):
            for c, val in enumerate(row_data):
                cell = tbl.cell(r + 1, c)
                cell.text = val
                p = cell.text_frame.paragraphs[0]
                p.runs[0].font.size = Pt(10)
                p.runs[0].font.color.rgb = DARK_BG
                p.runs[0].font.name = "Calibri"
        break

# ── LIMITATIONS & NEXT STEPS (Text 29 → Text 28) ──
set_multiline(find_shape(sl, "Text 28"), [
    ("Does not do yet: Real-time extraction, automated retraining, multi-supplier overlap, cloud deployment", 10.5, False, DARK_BG),
    ("Data / scope: Synthetic data (540 days, 5 SKUs); real supplier notes may have more variability", 10.5, False, DARK_BG),
    ("Known failure modes: LLM API outage degrades to ERP-only features (graceful fallback); daily MAPE of ~23% is typical for SKU-level forecasting", 10.5, False, DARK_BG),
    ("Next step: Hyperparameter tuning (Optuna), real-time streaming extraction, automated retraining on PSI drift alerts", 10.5, False, DARK_BG),
])

# Footer
set_text(find_shape(sl, "Text 29"),
         "EXL \u00b7 AI Services Capstone \u00b7 Demand Signal Feature Store \u2014 One-page summary",
         font_size=9, color=HINT)

# ── Save ──
out_path = r"C:\Users\LKDJ\Documents\DemandSignalFeatureStore\ONE_PAGE_SUMMARY.pptx"
prs.save(out_path)
print(f"Saved: {out_path}")
