"""Fill the official Capstone Pitch Deck template with project content."""

import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from copy import deepcopy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

TEMPLATE = r"C:\Users\LKDJ\Downloads\Capstone_Pitch_Deck_Template.pptx"
prs = Presentation(TEMPLATE)

# ── Colors from template ──
EXL_ORANGE   = RGBColor(0xFB, 0x4E, 0x0B)
EXL_PEACH    = RGBColor(0xFF, 0x8A, 0x5B)
DARK_BG      = RGBColor(0x1C, 0x1C, 0x1C)
WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
CREAM        = RGBColor(0xF3, 0xE7, 0xDF)
MUTED        = RGBColor(0x6E, 0x6A, 0x66)
HINT         = RGBColor(0x9B, 0x95, 0x8F)
DARK_DEMO    = RGBColor(0x2A, 0x2A, 0x2A)

def find_shape(slide, name):
    for s in slide.shapes:
        if s.name == name:
            return s
    return None

def set_text(shape, text, font_size=None, bold=None, color=None, font_name=None):
    """Replace text in shape, preserving existing formatting when params are None."""
    if shape is None:
        return
    tf = shape.text_frame
    # Clear all but first paragraph
    while len(tf.paragraphs) > 1:
        p = tf.paragraphs[-1]._element
        p.getparent().remove(p)
    tf.paragraphs[0].clear()
    run = tf.paragraphs[0].add_run()
    run.text = text
    if font_size:
        run.font.size = Pt(font_size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = color
    if font_name:
        run.font.name = font_name

def set_multiline(shape, lines, default_size=12, default_color=None, default_bold=False, default_name="Calibri"):
    """Set multiple paragraphs in a shape.
    lines = [(text, size, bold, color), ...]
    """
    if shape is None:
        return
    tf = shape.text_frame
    # Remove existing paragraphs
    for p in list(tf.paragraphs):
        p._element.getparent().remove(p._element)
    # Add new paragraphs
    from pptx.oxml import parse_xml
    from lxml import etree
    txBody = tf._txBody
    for i, (text, size, bold, color) in enumerate(lines):
        p_xml = parse_xml(
            f'<a:p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:r><a:t></a:t></a:r></a:p>'
        )
        txBody.append(p_xml)
        p = tf.paragraphs[i]
        p.runs[0].text = text
        p.runs[0].font.size = Pt(size or default_size)
        p.runs[0].font.bold = bold if bold is not None else default_bold
        p.runs[0].font.color.rgb = color or default_color or DARK_BG
        p.runs[0].font.name = default_name


def add_text_box(slide, left, top, width, height, text, size=12, bold=False,
                 color=DARK_BG, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                      Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_bullet_box(slide, left, top, width, height, items, size=12,
                   color=DARK_BG, spacing=Pt(3), font_name="Calibri"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                      Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if isinstance(item, tuple):
            txt, bld, clr = item
        else:
            txt, bld, clr = item, False, color
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = txt
        p.font.size = Pt(size)
        p.font.bold = bld
        p.font.color.rgb = clr
        p.font.name = font_name
        p.space_after = spacing
    return txBox


def add_rounded_box(slide, left, top, width, height, text, fill_color,
                    font_color=WHITE, font_size=12, bold=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                    Inches(left), Inches(top),
                                    Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = font_color
    p.font.name = "Calibri"
    p.alignment = PP_ALIGN.CENTER
    return shape


# ════════════════════════════════════════════════════════════════
# SLIDE 1 — TITLE (fill placeholders)
# ════════════════════════════════════════════════════════════════
sl = prs.slides[0]
set_text(find_shape(sl, "Text 3"), "Demand Signal Feature Store",
         font_size=54, bold=True, color=WHITE, font_name="Cambria")
set_text(find_shape(sl, "Text 4"),
         "LLM-extracted supplier signals improve demand forecasting by 6.6%",
         font_size=18, color=CREAM, font_name="Calibri")
set_text(find_shape(sl, "Text 6"), "Track 2", font_size=14, color=WHITE)
set_text(find_shape(sl, "Text 8"), "Supply Chain / Retail", font_size=14, color=WHITE)
set_text(find_shape(sl, "Text 10"), "Damie", font_size=14, color=WHITE)
set_text(find_shape(sl, "Text 11"),
         "Presented by: Damie   \u00b7   Date: 29 Sep 2026",
         font_size=12, color=CREAM)

# ════════════════════════════════════════════════════════════════
# SLIDE 2 — PROBLEM & BUSINESS CASE
# ════════════════════════════════════════════════════════════════
sl = prs.slides[1]
set_text(find_shape(sl, "Text 1"),
         "ERP forecasts miss early supplier delay signals",
         font_size=30, bold=True, color=DARK_BG, font_name="Cambria")

# "The problem" content (Text 6)
set_multiline(find_shape(sl, "Text 6"), [
    ("\u2022  ERP systems track structured data: orders, inventory, shipments", 12, False, DARK_BG),
    ("\u2022  Supplier delay warnings arrive as free-text notes 5\u201312 days before ERP is updated", 12, False, DARK_BG),
    ("\u2022  During this blind spot, the demand forecast is based on stale assumptions", 12, False, DARK_BG),
    ("\u2022  Leads to inaccurate forecasts, excess safety stock, or stockouts", 12, False, DARK_BG),
])

# "Cost of the status quo" content (Text 9)
set_multiline(find_shape(sl, "Text 9"), [
    ("\u2022  Forecast errors of 25%+ MAPE when relying only on ERP data", 12, False, DARK_BG),
    ("\u2022  Inventory misallocation during supplier disruption windows", 12, False, DARK_BG),
    ("\u2022  Manual monitoring of supplier emails — not scalable", 12, False, DARK_BG),
    ("\u2022  No structured extraction from unstructured supplier communications", 12, False, DARK_BG),
])

# "Why now / why AI" content (Text 12)
set_multiline(find_shape(sl, "Text 12"), [
    ("\u2022  LLMs (GPT-4o) can extract structured fields from free-text supplier notes", 12, False, WHITE),
    ("\u2022  Feature stores enable systematic integration of LLM outputs with ERP data", 12, False, WHITE),
    ("\u2022  Drift detection (PSI) ensures extracted features remain reliable over time", 12, False, WHITE),
    ("\u2022  Result: automated, scalable early-warning pipeline for demand forecasting", 12, False, WHITE),
])

# Speaker guidance (Text 13)
set_text(find_shape(sl, "Text 13"),
         "Speaker guidance: explain the 5\u201312 day gap between note arrival and ERP update \u2014 this is the core value proposition.",
         font_size=10.5, color=HINT)

# ════════════════════════════════════════════════════════════════
# SLIDE 3 — SOLUTION ARCHITECTURE
# ════════════════════════════════════════════════════════════════
sl = prs.slides[2]
set_text(find_shape(sl, "Text 1"),
         "Feature store pipeline: notes + ERP \u2192 enhanced forecast",
         font_size=30, bold=True, color=DARK_BG, font_name="Cambria")

# Pipeline boxes (replacing the 5-box flow)
set_text(find_shape(sl, "Text 5"), "Supplier Notes\n(58 texts)", font_size=12, bold=True, color=WHITE)
set_text(find_shape(sl, "Text 8"), "GPT-4o\nExtraction", font_size=12, bold=True, color=WHITE)
set_text(find_shape(sl, "Text 11"), "Feature Store\n(Parquet)", font_size=12, bold=True, color=WHITE)
set_text(find_shape(sl, "Text 14"), "XGBoost\nRegression", font_size=12, bold=True, color=DARK_BG)
set_text(find_shape(sl, "Text 17"), "Next-Day\nForecast", font_size=12, bold=True, color=WHITE)

# Architecture guidance (Text 18)
set_text(find_shape(sl, "Text 18"),
         "Data flows left \u2192 right: 58 supplier notes are processed by GPT-4o into structured signals, merged with 15 ERP features in a Parquet feature store, then fed to XGBoost for next-day sales prediction.",
         font_size=10.5, color=HINT)

# "Key components & choices" (Text 21)
set_multiline(find_shape(sl, "Text 21"), [
    ("\u2022  LLM Extraction \u2014 GPT-4o extracts early_delay_flag, expected_delay_days, delay_reason from each note", 12, False, DARK_BG),
    ("\u2022  Feature Store \u2014 18 features in Parquet with lineage metadata and 24h freshness SLA", 12, False, DARK_BG),
    ("\u2022  Drift Monitoring \u2014 PSI/CSI tracks feature stability; alerts on distribution shift", 12, False, DARK_BG),
    ("\u2022  Leakage Prevention \u2014 chronological split, lag shifting, temporal windowing of note signals", 12, False, DARK_BG),
])

# "Specialization competencies" (Text 24)
set_multiline(find_shape(sl, "Text 24"), [
    ("\u2022  Generative AI \u2014 LLM-based structured extraction from unstructured text", 12, False, WHITE),
    ("\u2022  ML Engineering \u2014 XGBoost regression, feature engineering, chronological evaluation", 12, False, WHITE),
    ("\u2022  Data Engineering \u2014 Feature store design with Parquet storage and lineage tracking", 12, False, WHITE),
    ("\u2022  MLOps \u2014 PSI drift detection, freshness SLA monitoring, production-readiness patterns", 12, False, WHITE),
])

# ════════════════════════════════════════════════════════════════
# SLIDE 4 — LIVE DEMO
# ════════════════════════════════════════════════════════════════
sl = prs.slides[3]
set_text(find_shape(sl, "Text 1"),
         "End-to-end notebook: data \u2192 extraction \u2192 forecast",
         font_size=30, bold=True, color=WHITE, font_name="Cambria")

# Demo placeholder (Text 5)
set_text(find_shape(sl, "Text 5"),
         "\u25b6  Run notebooks/main.ipynb end-to-end in VS Code",
         font_size=18, bold=True, color=CREAM)

# Demo guidance (Text 6)
set_text(find_shape(sl, "Text 6"),
         "Run the notebook live. Show: (1) data generation, (2) LLM extraction with real API calls, (3) feature store creation, (4) model comparison. Have saved outputs as fallback.",
         font_size=11, color=HINT)

# Demo script (Text 9)
set_multiline(find_shape(sl, "Text 9"), [
    ("\u2022  Step 1: Generate synthetic data \u2014 show 2,700 daily records, 92 POs, 58 notes", 12, False, WHITE),
    ("\u2022  Step 2: Run LLM extraction \u2014 watch GPT-4o process notes, 48/58 flagged as delays", 12, False, WHITE),
    ("\u2022  Step 3: Build feature store \u2014 18 features saved to Parquet with lineage metadata", 12, False, WHITE),
    ("\u2022  Step 4: Train models \u2014 Baseline MAPE 25.01% vs Enhanced MAPE 23.36%", 12, False, WHITE),
    ("\u2022  Step 5: Show feature importance \u2014 early_delay_flag is #1", 12, False, WHITE),
    ("\u2022  Step 6: Drift detection \u2014 PSI confirms key features are stable", 12, False, WHITE),
], default_color=WHITE)

# ════════════════════════════════════════════════════════════════
# SLIDE 5 — RESULTS & EVALUATION
# ════════════════════════════════════════════════════════════════
sl = prs.slides[4]
set_text(find_shape(sl, "Text 1"),
         "6.6% MAPE improvement with note-derived features",
         font_size=30, bold=True, color=DARK_BG, font_name="Cambria")

# Quality metric
set_text(find_shape(sl, "Text 5"), "MAPE", font_size=13, bold=True, color=EXL_ORANGE)
set_text(find_shape(sl, "Text 6"), "-6.6%", font_size=34, bold=True, color=DARK_BG, font_name="Cambria")
set_text(find_shape(sl, "Text 7"), "25.01% \u2192 23.36%", font_size=12, color=MUTED)

# Latency metric
set_text(find_shape(sl, "Text 9"), "MAE", font_size=13, bold=True, color=EXL_ORANGE)
set_text(find_shape(sl, "Text 10"), "-9.5%", font_size=34, bold=True, color=DARK_BG, font_name="Cambria")
set_text(find_shape(sl, "Text 11"), "5.03 \u2192 4.55 units", font_size=12, color=MUTED)

# Cost metric
set_text(find_shape(sl, "Text 13"), "RMSE", font_size=13, bold=True, color=EXL_ORANGE)
set_text(find_shape(sl, "Text 14"), "-6.6%", font_size=34, bold=True, color=DARK_BG, font_name="Cambria")
set_text(find_shape(sl, "Text 15"), "6.80 \u2192 6.35 units", font_size=12, color=MUTED)

# Evaluation method (Text 18)
set_multiline(find_shape(sl, "Text 18"), [
    ("    Chronological Split \u2014 80% train / 20% test by date (no random shuffle)", 12, False, DARK_BG),
    ("    Same Hyperparameters \u2014 Both models use identical XGBoost config for fair comparison", 12, False, DARK_BG),
    ("    Feature Importance \u2014 early_delay_flag is the #1 feature in the enhanced model", 12, False, DARK_BG),
    ("    Synthetic Data \u2014 540 days, 5 SKUs, 4 suppliers with 35\u201355% reliability", 12, False, DARK_BG),
])

# Safety & guardrails → Monitoring & drift (Text 21)
set_text(find_shape(sl, "Text 20"), "Monitoring & drift detection",
         font_size=14, bold=True, color=EXL_PEACH)
set_multiline(find_shape(sl, "Text 21"), [
    ("    PSI Drift Detection \u2014 early_delay_flag PSI = 0.07 (stable), delay_reason PSI = 1.06 (excluded)", 12, False, WHITE),
    ("    Freshness SLA \u2014 24-hour check on all feature timestamps before inference", 12, False, WHITE),
    ("    Lineage Tracking \u2014 Every extraction tagged with timestamp, model version, source", 12, False, WHITE),
    ("    Feature Store \u2014 Parquet storage with CSV lineage metadata for full traceability", 12, False, WHITE),
])

# ════════════════════════════════════════════════════════════════
# SLIDE 6 — LIMITATIONS & NEXT STEPS
# ════════════════════════════════════════════════════════════════
sl = prs.slides[5]
set_text(find_shape(sl, "Text 1"),
         "Honest constraints and the roadmap forward",
         font_size=30, bold=True, color=DARK_BG, font_name="Cambria")

# Limitations (Text 6)
set_multiline(find_shape(sl, "Text 6"), [
    ("\u2022  Synthetic data only \u2014 real supplier notes may have more variability", 12, False, DARK_BG),
    ("\u2022  Single notebook, not a deployed service \u2014 no API or scheduling", 12, False, DARK_BG),
    ("\u2022  No hyperparameter tuning \u2014 MAPE could be lower with optimization", 12, False, DARK_BG),
    ("\u2022  One supplier per SKU \u2014 no multi-supplier overlap handling", 12, False, DARK_BG),
])

# Trade-offs (Text 9)
set_multiline(find_shape(sl, "Text 9"), [
    ("\u2022  Excluded delay_reason from model \u2014 traded semantic richness for stability (PSI confirmed)", 12, False, DARK_BG),
    ("\u2022  XGBoost over linear regression \u2014 traded interpretability for non-linear accuracy", 12, False, DARK_BG),
    ("\u2022  Daily granularity \u2014 traded lower MAPE (weekly) for more actionable daily predictions", 12, False, DARK_BG),
])

# Next steps (Text 12)
set_multiline(find_shape(sl, "Text 12"), [
    ("\u2022  Hyperparameter tuning (Optuna/GridSearchCV) for further MAPE reduction", 12, False, WHITE),
    ("\u2022  Azure deployment \u2014 Blob Storage for features, Azure ML for inference", 12, False, WHITE),
    ("\u2022  Real-time extraction \u2014 streaming LLM as notes arrive", 12, False, WHITE),
    ("\u2022  Automated retraining triggered by PSI drift alerts", 12, False, WHITE),
])

# Thank you line (Text 13)
set_text(find_shape(sl, "Text 13"),
         "Thank you \u2014 questions welcome.",
         font_size=13, bold=True, color=EXL_ORANGE)

# ── Update page numbers ──
for i, sl in enumerate(prs.slides):
    s = find_shape(sl, "Text 15") or find_shape(sl, "Text 11" if i == 3 else "NONE")
    # Page numbers are in the last text shape with "X / 6" pattern
    for shape in sl.shapes:
        if hasattr(shape, 'text') and '/ 6' in shape.text:
            set_text(shape, f"{i+1} / 6", font_size=9, color=MUTED if i != 3 else CREAM)

# ── Save ──
out_path = r"C:\Users\LKDJ\Documents\DemandSignalFeatureStore\PITCH_DECK.pptx"
prs.save(out_path)
print(f"Saved: {out_path}")
