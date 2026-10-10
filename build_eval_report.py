"""Build the Evaluation Report as a formatted Word document."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

doc = Document()

section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

EXL_ORANGE = RGBColor(0xFB, 0x4D, 0x0A)
EXL_SLATE  = RGBColor(0x2E, 0x36, 0x43)
EXL_BLUE   = RGBColor(0x00, 0x50, 0x71)
EXL_GRAY   = RGBColor(0x6D, 0x72, 0x7B)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GRAY  = RGBColor(0x33, 0x33, 0x33)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.font.color.rgb = DARK_GRAY
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.15

for level in range(1, 3):
    h = doc.styles[f'Heading {level}']
    h.font.name = 'Calibri'
    if level == 1:
        h.font.size = Pt(22)
        h.font.color.rgb = EXL_ORANGE
        h.font.bold = True
        h.paragraph_format.space_before = Pt(24)
        h.paragraph_format.space_after = Pt(8)
    elif level == 2:
        h.font.size = Pt(16)
        h.font.color.rgb = EXL_BLUE
        h.font.bold = True
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)


def add_para(text, bold=False, size=None, color=None, align=None, space_after=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if bold: run.bold = True
    if size: run.font.size = Pt(size)
    if color: run.font.color.rgb = color
    if align: p.alignment = align
    if space_after is not None: p.paragraph_format.space_after = Pt(space_after)
    return p


def add_bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    p.text = text
    p.paragraph_format.left_indent = Cm(1.2)
    return p


def add_table(headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = h
        p = cell.paragraphs[0]
        p.runs[0].bold = True
        p.runs[0].font.color.rgb = WHITE
        p.runs[0].font.size = Pt(10)
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="005071"/>')
        cell._element.get_or_add_tcPr().append(shading)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.rows[r + 1].cells[c]
            cell.text = str(val)
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(10)
            bg = "F5F5F5" if r % 2 == 0 else "FFFFFF"
            shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg}"/>')
            cell._element.get_or_add_tcPr().append(shading)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def add_callout(text, bg_color="E6F6FB"):
    p = doc.add_paragraph()
    pPr = p._element.get_or_add_pPr()
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
    pPr.append(shading)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.color.rgb = EXL_BLUE
    run.italic = True
    return p


# ── COVER ──
doc.add_paragraph()
doc.add_paragraph()
doc.add_paragraph()

add_para("Evaluation Report", bold=True, size=32, color=EXL_ORANGE,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
add_para("Demand Signal Feature Store", bold=True, size=20, color=EXL_SLATE,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

doc.add_paragraph()

add_para("Capstone Project  |  AI Engineer Program  |  Track 1", size=12, color=EXL_GRAY,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
add_para("Damini Thandele  |  October 2026", size=12, color=EXL_GRAY,
         align=WD_ALIGN_PARAGRAPH.CENTER)

doc.add_page_break()

# ── 1. OBJECTIVE ──
doc.add_heading("1. Objective", level=1)

doc.add_paragraph(
    "Measure whether LLM-extracted supplier delay signals improve next-day demand forecasting "
    "accuracy compared to an ERP-only baseline. The primary metric is MAPE."
)

# ── 2. EXPERIMENTAL SETUP ──
doc.add_heading("2. Experimental Setup", level=1)

add_table(
    ["Parameter", "Value"],
    [
        ["Algorithm", "XGBoost Regressor"],
        ["Hyperparameters", "n_estimators=200, max_depth=5, learning_rate=0.1, random_state=42"],
        ["Train/Test Split", "Chronological 80/20 (first ~14 months train, last ~4 months test)"],
        ["Baseline Features", "15 ERP features (sales lags, inventory, PO timing, calendar)"],
        ["Enhanced Features", "15 ERP + 2 note-derived (early_delay_flag, expected_delay_days)"],
        ["Target Variable", "target_next_day_sales"],
        ["Data", "Synthetic: 2,700 daily records, 5 SKUs, 4 suppliers, 540 days"],
    ],
    col_widths=[4, 12]
)

doc.add_paragraph(
    "Both models use identical hyperparameters so any accuracy difference comes purely from "
    "the additional features."
)

# ── 3. LLM EXTRACTION QUALITY ──
doc.add_heading("3. LLM Extraction Quality", level=1)

doc.add_paragraph(
    "The LLM (GPT-4o) extracts structured fields from 58 supplier notes. Ground truth comes "
    "from the data generator’s note_type label (delay_warning vs. routine)."
)

add_table(
    ["Metric", "Score", "Meaning"],
    [
        ["Precision", "100%", "Of notes flagged as delays, all were actual delays"],
        ["Recall", "100%", "All actual delay notes were correctly identified"],
        ["F1 Score", "100%", "Harmonic mean of precision and recall"],
        ["Accuracy", "100%", "Overall correct classification rate (58/58)"],
        ["Delay Days MAE", "~1.5 days", "Average error in predicted delay duration vs. actual"],
    ],
    col_widths=[3, 2, 8.5]
)

add_callout(
    "Scores are near-perfect on synthetic data. Real-world supplier notes with more variability "
    "would likely produce lower scores — this is a known limitation.",
    bg_color="FFF3E0"
)

doc.add_page_break()

# ── 4. FORECAST RESULTS ──
doc.add_heading("4. Forecast Results", level=1)

add_table(
    ["Metric", "Baseline (ERP only)", "Enhanced (ERP + Notes)", "Change"],
    [
        ["MAPE", "25.01%", "23.36%", "-6.6% (improved)"],
        ["MAE", "5.03 units", "4.55 units", "-9.5% (improved)"],
        ["RMSE", "6.80 units", "6.35 units", "-6.6% (improved)"],
    ],
    col_widths=[3, 3.5, 3.5, 3.5]
)

add_callout(
    "MAPE improvement: (23.36 \u2212 25.01) / 25.01 \u00d7 100 = \u22126.6%. "
    "All three metrics improved consistently.",
    bg_color="E8F5E9"
)

# ── 5. FEATURE IMPORTANCE ──
doc.add_heading("5. Feature Importance (Top 5)", level=1)

add_table(
    ["Rank", "Feature", "Source", "Role"],
    [
        ["1", "early_delay_flag", "Supplier Note (LLM)", "Binary delay warning \u2014 available before ERP update"],
        ["2", "sales_lag_1", "ERP", "Previous day\u2019s sales"],
        ["3", "rolling_avg_7", "ERP", "7-day smoothed demand trend"],
        ["4", "on_hand_inventory", "ERP", "Current warehouse stock"],
        ["5", "inventory_cover", "ERP", "Days of stock remaining"],
    ],
    col_widths=[1.2, 3.5, 3, 5.8]
)

doc.add_paragraph(
    "The LLM-extracted early_delay_flag is the #1 most important feature, validating that "
    "unstructured supplier signals carry unique predictive value."
)

# ── 6. DRIFT DETECTION ──
doc.add_heading("6. Drift Detection (PSI)", level=1)

add_table(
    ["Feature", "PSI", "Status"],
    [
        ["sales_lag_1", "0.03", "Stable"],
        ["on_hand_inventory", "0.01", "Stable"],
        ["early_delay_flag", "0.07", "Stable"],
        ["inventory_cover", "0.12", "Moderate"],
        ["rolling_avg_7", "0.24", "Significant"],
        ["expected_delay_days", "0.63", "Significant"],
        ["historical_avg_supplier_delay", "0.78", "Significant"],
        ["delay_reason_encoded", "1.06", "Significant"],
    ],
    col_widths=[5, 2, 3]
)

doc.add_paragraph(
    "Thresholds: <0.1 stable, 0.1\u20130.2 moderate, \u22650.2 significant. "
    "The key note feature (early_delay_flag, PSI = 0.07) is stable. "
    "delay_reason_encoded (PSI = 1.06) confirms the decision to exclude it from the model."
)

# ── 7. CONCLUSION ──
doc.add_heading("7. Conclusion", level=1)

doc.add_paragraph(
    "Adding two LLM-extracted features to 15 ERP features reduces MAPE from 25.01% to 23.36% "
    "(\u22126.6%). The improvement is driven by early warning signals available 5\u201312 days before the "
    "ERP system is updated. The key feature is stable (PSI = 0.07), supporting production confidence."
)

doc.add_heading("Limitations", level=2)

add_bullet("Synthetic data \u2014 real supplier notes may have more variability")
add_bullet("No hyperparameter tuning \u2014 MAPE could be further reduced with optimization")
add_bullet("Single supplier per SKU \u2014 no multi-supplier overlap handling")
add_bullet("Observed sales underestimates true demand during stockouts")

# ── Save ──
out_path = r"C:\Users\LKDJ\Documents\DemandSignalFeatureStore\EVALUATION_REPORT.docx"
doc.save(out_path)
print(f"Saved: {out_path}")
